# -*- coding: utf-8 -*-
"""M1 装配：把内容填进官方决赛模板，输出主文档 docx。

内容来源：
- build_finals_content.CONTENT（一~五节）、TEAM、LINKS、BAILIAN_ROWS
- build_finals_appendix.SECTION6（第六节）与 APPENDIX（附录 A~F）

用法：python build_finals_docx.py [--out 文件名.docx]
"""
from __future__ import annotations

import argparse
import os
import sys

from docx import Document
from docx.shared import Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_finals_appendix import APPENDIX, SECTION6  # noqa: E402
from build_finals_content import BAILIAN_ROWS, CONTENT, LINKS, TEAM  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "hackathon-决赛入围定稿作品提交模板-天池版.docx")
DEFAULT_OUT = os.path.join(HERE, "更新世界的锋芒_SoundInsight_决赛入围定稿作品.docx")

PLAN = dict(CONTENT)
PLAN["六、项目开发及阶段成果说明"] = SECTION6


def set_cell(cell, text: str) -> None:
    cell.text = ""
    cell.paragraphs[0].add_run(text)


def fill_tables(doc) -> None:
    for t in doc.tables:
        header = " ".join(c.text for c in t.rows[0].cells)
        is_bailian = "模型标识" in header
        is_links = ("链接类型" in header) or ("链接地址" in header)
        if is_bailian:
            idx = 1
            for label, usage, link in BAILIAN_ROWS:
                row = t.rows[idx] if idx < len(t.rows) else t.add_row()
                if len(row.cells) >= 3:
                    set_cell(row.cells[0], label)
                    set_cell(row.cells[1], usage)
                    set_cell(row.cells[2], link)
                idx += 1
            continue
        # 全体成员表：表头下一行填队长信息（数据行首格为空，故按"表头行号 + 1"定位）
        for ri, row in enumerate(t.rows):
            cells = row.cells
            if len(cells) >= 6 and cells[0].text.strip() == "姓名" and ri + 1 < len(t.rows):
                drow = t.rows[ri + 1].cells
                if len(drow) >= 6 and not drow[0].text.strip():
                    set_cell(drow[0], TEAM["队长姓名"])
                    set_cell(drow[1], TEAM["联系电话"])
                    set_cell(drow[2], TEAM["联系邮箱"])
                    set_cell(drow[3], "个人团队")
                    set_cell(drow[4], "个人开发者")
                    set_cell(drow[5], "队长")
                break
        for row in t.rows:
            cells = row.cells
            if not cells or len(cells) < 2:
                continue
            key = cells[0].text.strip()
            # 表格里的示例/提示行：清空（如"（请将各链接粘贴到对应行）💡 建议…"）
            if key.startswith("（") or "请将各链接" in key:
                for c in cells:
                    set_cell(c, "")
                continue
            if key in TEAM and not cells[1].text.strip():
                set_cell(cells[1], TEAM[key])
            elif key == "参赛场景":
                set_cell(cells[1], "AI市场洞察")
            elif key == "方案名称":
                set_cell(cells[1], "SoundInsight")
            elif key.startswith("一句话定义"):
                set_cell(cells[1], "为跨境电商耳机卖家提供音质差评自动识别与五类问题归因的"
                                   "一站式洞察工具，把人工数小时的差评梳理压缩到分钟级；"
                                   "产品推理 100% 本地。")
            elif is_links:
                for k, v in LINKS.items():
                    if key.startswith(k):
                        set_cell(cells[1], v)


def split_blocks(lines: list[str]):
    """把内容行切成块：[("p", 文本)] 与 [("table", 行列表)]。

    Markdown 管道表格（连续的 `| ... |` 行，第二行为分隔行）转成真正的 Word 表格——
    否则在 docx 里就是一串带竖线的纯文本，观感与可读性都很差。
    """
    blocks, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("|") and ln.count("|") >= 2:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") and c for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        blocks.append(("p", ln))
        i += 1
    return blocks


def is_guidance(text: str) -> bool:
    """判定模板引导语：整段以（ 开头、）结尾，或以"（例：""（请""（如"起。

    这类段落必须在填充后删除——留着会让评委以为文档没填完（2026-10-01 实测残留 10 段）。
    """
    t = text.strip()
    if not t:
        return False
    if t.startswith(("（例：", "（请", "（如", "（可列")):
        return True
    return t.startswith("（") and t.endswith("）") and len(t) < 240


def strip_guidance(doc) -> int:
    """删除正文层的模板引导语段落 + 4.2 的空编号占位（1. 2. 3. 4.）。"""
    removed = 0
    numbered = {"1.", "2.", "3.", "4.", "5."}
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if is_guidance(t) or t in numbered:
            p._element.getparent().remove(p._element)
            removed += 1
    return removed


def polish_cover_and_checklist(doc) -> tuple[bool, int]:
    """封面换成"定稿作品"并加方案行；勾选第七节提交物清单（我们四件齐备）。

    返回 (是否改了封面, 勾选数)。模板原文"提交模板"字样若留在提交物里，会显得没定稿。
    """
    cover_done, ticked = False, 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("AI+跨境黑客松巅峰赛") and "模板" in t:
            for run in p.runs:
                run.text = ""
            if p.runs:
                p.runs[0].text = "AI+跨境黑客松巅峰赛 · 决赛入围定稿作品"
            else:
                p.add_run("AI+跨境黑客松巅峰赛 · 决赛入围定稿作品")
            # 在其后插入方案信息行
            new_p = p.insert_paragraph_before("")
            p._p.addnext(new_p._p)
            new_p.text = ("SoundInsight —— 蓝牙耳机音质差评智能归因系统　｜　"
                          "团队：更新世界的锋芒　｜　参赛场景：AI 市场洞察")
            for run in new_p.runs:
                run.bold = True
                run.font.size = Pt(12)
            cover_done = True
        elif t.startswith("☐"):
            for run in p.runs:
                run.text = run.text.replace("☐", "☑")
            if "☑" not in p.text:      # 无 run 的情况下直接改段首
                p.text = "☑" + t[1:]
            ticked += 1
    return cover_done, ticked


def fix_participation_section(doc) -> tuple[int, int, int]:
    """清理"二、参赛信息"的重复与占位：

    - 模板结构是「表（参赛场景／方案名称） + 一行『一句话定义（…）：』」；
      我此前又在表格前插入了三行同内容段落 → **重复**（评委能看到两遍）。
    - 本函数：① 删除重复的"参赛场景：/方案名称：/一句话定义："段落；
      ② 把模板那行占位改写成真正的定义行。
    """
    removed = rewritten = 0
    definition = ("为跨境电商耳机卖家提供音质差评自动识别与五类问题归因的一站式洞察工具，"
                  "把人工数小时的差评梳理压缩到分钟级；产品推理 100% 本地，卖家数据不出境。")
    for p in list(doc.paragraphs):
        t = p.text.strip()
        if t.startswith(("参赛场景：", "方案名称：", "一句话定义：")):
            p._element.getparent().remove(p._element)
            removed += 1
        elif t.startswith("一句话定义（"):
            for run in p.runs:
                run.text = ""
            if p.runs:
                p.runs[0].text = "一句话定义：" + definition
            else:
                p.add_run("一句话定义：" + definition)
            rewritten += 1
    # 把「适用范围与价值边界」段移到定义之后（评委应先读到价值，再读到边界）
    def_p = boundary_p = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("一句话定义：") and def_p is None:
            def_p = p
        elif ("价值边界" in t) and boundary_p is None:
            boundary_p = p
    if def_p is not None and boundary_p is not None:
        boundary_p._p.getparent().remove(boundary_p._p)
        def_p._p.addnext(boundary_p._p)
        moved = 1
    else:
        moved = 0
    return removed, rewritten, moved


def apply_inline_format(doc) -> tuple[int, int]:
    """把 Markdown 行内标记转为 Word 格式：`**粗体**` → bold run；反引号 → 去掉（保留文本）。

    只处理含标记的段落，避免影响模板自带的样式与其它 run 级格式。
    """
    n_bold = n_code = 0

    def fix_paragraph(p) -> None:
        nonlocal n_bold, n_code
        t = p.text or ""
        if "**" not in t and "`" not in t:
            return
        segs = []
        for k, part in enumerate(t.split("**")):
            if part == "":
                continue
            if "`" in part:
                n_code += part.count("`")
                part = part.replace("`", "")
            segs.append((part, k % 2 == 1))
        if not segs:
            return
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        for text, bold in segs:
            run = p.add_run(text)
            run.bold = bold
        n_bold += 1

    for p in doc.paragraphs:
        fix_paragraph(p)
    for tb in doc.tables:
        for row in tb.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fix_paragraph(p)
    return n_bold, n_code

def insert_after(doc, paragraph, lines: list[str]) -> int:
    """在段落之后依次插入内容（正文段落或 Word 表格），返回插入的块数。"""
    anchor = paragraph._p
    n = 0
    for kind, payload in split_blocks(lines):
        if kind == "p":
            new_p = paragraph.insert_paragraph_before("")
            anchor.addnext(new_p._p)
            new_p.text = payload
            # 附录标题加粗放大，便于评委跳读
            if payload.startswith("附录 ") or payload.startswith("模型调用与边界"):
                for run in new_p.runs:
                    run.bold = True
                    run.font.size = Pt(13)
            anchor = new_p._p
        else:
            rows, cols = len(payload), max(len(r) for r in payload)
            tbl = doc.add_table(rows=rows, cols=cols)
            try:
                tbl.style = "Table Grid"
            except KeyError:
                pass
            for ri, row in enumerate(payload):
                for ci in range(cols):
                    text = row[ci] if ci < len(row) else ""
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(text)
                    if ri == 0:
                        run.bold = True
                    run.font.size = Pt(9)
            anchor.addnext(tbl._tbl)      # 表格移到锚点之后
            anchor = tbl._tbl
        n += 1
    return n


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=DEFAULT_OUT)
    args = ap.parse_args()
    if not os.path.isfile(TEMPLATE):
        print(f"[FAIL] 缺少模板 {os.path.basename(TEMPLATE)}")
        return 1

    doc = Document(TEMPLATE)
    print(f"模板：段落 {len(doc.paragraphs)}、表格 {len(doc.tables)}")
    n_guide = strip_guidance(doc)
    print(f"删除模板引导语/占位段落：{n_guide}")
    cover, ticked = polish_cover_and_checklist(doc)
    print(f"封面改为定稿作品：{'是' if cover else '否'}｜提交物清单勾选：{ticked} 项")
    inserted, matched = 0, []
    for para in list(doc.paragraphs):
        text = para.text.strip()
        for key, lines in PLAN.items():
            if text == key or text.startswith(key):
                inserted += insert_after(doc, para, lines)
                matched.append(key)
                break
    print(f"已填充小节 {len(matched)}：{'、'.join(matched)}")
    print(f"插入内容块 {inserted}")
    rm, rw, mv = fix_participation_section(doc)
    print(f"参赛信息节：删除重复段落 {rm}、改写占位行 {rw}")
    # 附录：同样支持 Markdown 表格 → Word 表格
    app_blocks = split_blocks(APPENDIX)
    n_tbl = 0
    for kind, payload in app_blocks:
        if kind == "p":
            p = doc.add_paragraph(payload)
            if payload.startswith("附录 ") or payload.startswith("模型调用与边界"):
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(13)
        else:
            rows, cols = len(payload), max(len(r) for r in payload)
            tbl = doc.add_table(rows=rows, cols=cols)
            try:
                tbl.style = "Table Grid"
            except KeyError:
                pass
            for ri, row in enumerate(payload):
                for ci in range(cols):
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    run = cell.paragraphs[0].add_run(row[ci] if ci < len(row) else "")
                    if ri == 0:
                        run.bold = True
                    run.font.size = Pt(9)
            n_tbl += 1
    print(f"追加附录：{len(app_blocks)} 块（含 {n_tbl} 个表格）")
    fill_tables(doc)
    n_bold, n_code = apply_inline_format(doc)
    print(f"行内格式化：粗体 {n_bold} 处、去反引号 {n_code} 处")
    doc.save(args.out)
    print(f"[写出] {os.path.basename(args.out)}（{os.path.getsize(args.out)/1024:.0f} KB）")

    check = Document(args.out)
    parts = [p.text for p in check.paragraphs]
    for t in check.tables:                       # 关键数字现在可能在表格里
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    all_text = "\n".join(parts)
    must = ["0.7220", "0.7811", "0.8273", "已知表述勘误与口径演进",
            "置信度档位与建议动作", "qwen3.7-plus", "deepseek-chat", "0.5333", "68.3%"]
    for m in must:
        print(f"  {'OK  ' if m in all_text else 'MISS'} 含「{m}」")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
