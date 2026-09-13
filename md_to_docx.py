# -*- coding: utf-8 -*-
# md_to_docx.py —— 把复赛主文档（Markdown）转成真正的 Word 文档（.docx）
#
# 为什么需要它：把 .md 直接改扩展名成 .pdf/.docx 并不是转换，WPS/Word 会打不开。
# 本脚本生成合规的 .docx：WPS 可直接打开、编辑，并可一键"输出为PDF"。
#
# 用法：
#   python md_to_docx.py competition_v4.md "更新世界的锋芒_SoundInsight_复赛作品.docx"
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))

CJK = "等线"
LATIN = "Calibri"


def style_run(run, size=None, bold=None, mono=False):
    """设置字体（含东亚字体），否则中文可能显示为方框或回退字体。"""
    name = "Consolas" if mono else LATIN
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(par, text, size=10.5, base_bold=False):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = par.add_run(piece[2:-2])
            style_run(r, size=size, bold=True)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            style_run(r, size=size - 0.5, bold=base_bold, mono=True)
        else:
            r = par.add_run(piece)
            style_run(r, size=size, bold=base_bold)


def is_table_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def split_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def convert(md_path, out_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 页面默认字体
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    i = 0
    n_table = n_head = n_para = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            r = p.add_run("—" * 20)
            style_run(r, size=9)
            r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            h = doc.add_heading(level=min(level, 4))
            add_inline(h, text, size={1: 18, 2: 15, 3: 12.5, 4: 11.5}[min(level, 4)],
                       base_bold=True)
            n_head += 1
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) \
                and is_table_sep(lines[i + 1]):
            header = split_row(stripped)
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append(split_row(lines[j]))
                j += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for k, cell in enumerate(header):
                par = table.rows[0].cells[k].paragraphs[0]
                add_inline(par, cell, size=10, base_bold=True)
            for row in rows:
                cells = table.add_row().cells
                for k in range(len(header)):
                    txt = row[k] if k < len(row) else ""
                    add_inline(cells[k].paragraphs[0], txt, size=10)
            n_table += 1
            doc.add_paragraph()
            i = j
            continue

        # 列表（有序 / 无序）
        m = re.match(r"^([-*])\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(2))
            n_para += 1
            i += 1
            continue
        m = re.match(r"^(\d+)[.、)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            n_para += 1
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_inline(p, stripped)
        n_para += 1
        i += 1

    doc.save(out_path)
    print(f"段落 {n_para} ｜ 标题 {n_head} ｜ 表格 {n_table}")
    print(f"已生成 -> {out_path}（{os.path.getsize(out_path):,} B）")


def main():
    md = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HERE, "competition_v4.md")
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        HERE, "更新世界的锋芒_SoundInsight_复赛作品.docx")
    convert(md, out)


if __name__ == "__main__":
    main()
