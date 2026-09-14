# -*- coding: utf-8 -*-
# 本脚本是 SoundInsight 一键洞察 Agent：输入评论 CSV，自动完成音质差评识别、
# 问题归因与统计，输出洞察报告（markdown/excel）并打印摘要。一条命令即可运行。
# 支持 --lang en 英文报告；非英文评论经 predict_core.is_unsupported 跳过并计数。
import argparse
import json
import os
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import torch
from transformers import (DistilBertForSequenceClassification,
                          DistilBertTokenizer)

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import report_builder as rb  # noqa: E402
from predict_core import is_unsupported  # noqa: E402

with open(os.path.join(HERE, "config.json"), encoding="utf-8") as f:
    CFG = json.load(f)
BIN_MODEL = os.path.join(HERE, CFG["bin_model_dir"])
ML_MODEL = os.path.join(HERE, CFG["multi_label_dir"])
THR_FILE = os.path.join(HERE, CFG["threshold_file"])
ISSUE_FILE = os.path.join(HERE, CFG["issue_labels_file"])
BATCH = int(CFG["batch_size"])
MAX_LEN = int(CFG["max_len"])


def load():
    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    tok = DistilBertTokenizer.from_pretrained(BIN_MODEL)
    bin_model = DistilBertForSequenceClassification.from_pretrained(
        BIN_MODEL).to(device)
    bin_model.eval()
    with open(THR_FILE, encoding="utf-8") as f:
        thr = float(json.load(f)["threshold"])
    ml_model = DistilBertForSequenceClassification.from_pretrained(
        ML_MODEL).to(device)
    ml_model.eval()
    with open(ISSUE_FILE, encoding="utf-8") as f:
        issue_info = json.load(f)
    return device, tok, bin_model, thr, ml_model, issue_info


def predict_probs(model, tok, texts, device):
    probs = []
    with torch.no_grad():
        for b in range(0, len(texts), BATCH):
            enc = tok(texts[b:b + BATCH], padding=True, truncation=True,
                      max_length=MAX_LEN, return_tensors="pt")
            enc = {k: v.to(device) for k, v in enc.items()}
            logits = model(**enc).logits
            probs.append(torch.softmax(logits, -1)[:, 1].cpu().numpy())
    return np.concatenate(probs)


def predict_multilabel(model, tok, texts, device):
    out = []
    with torch.no_grad():
        for b in range(0, len(texts), BATCH):
            enc = tok(texts[b:b + BATCH], padding=True, truncation=True,
                      max_length=MAX_LEN, return_tensors="pt")
            enc = {k: v.to(device) for k, v in enc.items()}
            out.append(torch.sigmoid(model(**enc).logits).cpu().numpy())
    return np.vstack(out)


def analyze(csv_path: str, export: bool = False, lang: str = "zh") -> str:
    device, tok, bin_model, thr, ml_model, issue_info = load()
    df = pd.read_csv(csv_path, encoding="utf-8")
    text_col = next((c for c in ("text", "reviewText", "review")
                     if c in df.columns), df.columns[0])
    df["_text"] = df[text_col].astype(str)
    n_unsup = int(df["_text"].map(is_unsupported).sum())
    df = df[~df["_text"].map(is_unsupported)].reset_index(drop=True)
    texts = df["_text"].tolist()
    print(f"分析 {len(texts)} 条评论 (device={device})，"
          f"跳过非英文 {n_unsup} 条")

    probs = predict_probs(bin_model, tok, texts, device)
    df["sound_negative_prob"] = probs
    df["is_sound_negative"] = (probs >= thr).astype(int)

    n = len(df)
    n_neg = int(df["is_sound_negative"].sum())
    neg_idx = df[df["is_sound_negative"] == 1].index.tolist()

    issue_cols = issue_info["columns"]
    issue_names = issue_info["names"]
    issue_counts = {k: 0 for k in issue_names}
    ml_probs_by_idx = {}
    if neg_idx:
        neg_texts = [texts[i] for i in neg_idx]
        ml_probs = predict_multilabel(ml_model, tok, neg_texts, device)
        ml_pred = (ml_probs >= 0.5).astype(int)
        for k, name in enumerate(issue_names):
            issue_counts[name] = int(ml_pred[:, k].sum())
        ml_probs_by_idx = {i: ml_probs[k] for k, i in enumerate(neg_idx)}
        df.loc[neg_idx, [c.replace("_llm", "") for c in issue_cols]] = ml_pred

    rating_col = next((c for c in ("rating", "overall", "score")
                       if c in df.columns), None)
    if rating_col is not None:
        avg_rating = float(pd.to_numeric(df[rating_col],
                                         errors="coerce").mean())
    else:
        avg_rating = float("nan")

    # 优先级规则与报告生成统一走 report_builder（与 Demo、在线版同一套口径）
    ranked, priority = rb.rank_issues(issue_counts)

    rate = (n_neg / n) if n else 0.0
    n_mid = int(((probs >= 0.5) & (probs < thr)).sum())
    examples = []
    for i in sorted(neg_idx, key=lambda j: -probs[j])[:5]:
        ip = {}
        if i in ml_probs_by_idx:
            ip = {name: float(ml_probs_by_idx[i][k])
                  for k, name in enumerate(issue_names)}
        examples.append({"text": texts[i], "prob": float(probs[i]),
                         "issue_probs": ip})

    src_name = os.path.basename(csv_path)
    report = rb.build_report(
        src_name=src_name, n_total=n + n_unsup, n_unsupported=n_unsup,
        n_valid=n, n_neg=n_neg, avg_rating=avg_rating,
        issue_counts=issue_counts, examples=examples, n_mid=n_mid, lang=lang)
    out_path = os.path.join(
        HERE,
        "insight_report_v2_en.md" if lang == "en" else "insight_report_v2.md")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(report)
    print(f"报告已保存 -> {out_path}")
    if export:
        export_excel(csv_path, n, rate, avg_rating, rb.verdict(rate), ranked,
                     priority, neg_idx, probs, texts)
    return report


def export_excel(csv_path, n, rate, avg_rating, verdict, ranked, priority,
                 neg_idx, probs, texts):
    """导出 Excel 版报告：总体概况 / 问题分布（高优先级标红）/ 典型案例。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "总体概况"
    rows1 = [["指标", "数值"],
             ["分析对象", os.path.basename(csv_path)],
             ["评论总数", n],
             ["音质差评数", f"{len(neg_idx)}（{rate:.2%}）"],
             ["平均评分", f"{avg_rating:.2f}"],
             ["结论", verdict]]
    for r in rows1:
        ws1.append(r)

    ws2 = wb.create_sheet("问题分布")
    ws2.append(["问题类别", "数量", "占比", "优先级"])
    total_issue = sum(c for _, c in ranked)
    red_font = Font(color="FF0000", bold=True)
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    for name, cnt in ranked:
        ws2.append([name, cnt,
                    f"{cnt / max(total_issue, 1):.1%}", priority[name]])
        if priority[name] == "高":
            for cell in ws2[ws2.max_row]:
                cell.font = red_font
                cell.fill = red_fill

    ws3 = wb.create_sheet("典型案例")
    ws3.append(["概率", "类别", "评论原文"])
    for i in sorted(neg_idx, key=lambda j: -probs[j])[:5]:
        ws3.append([f"{probs[i]:.1%}", "音质负面", texts[i][:500]])

    xlsx_path = os.path.join(HERE, "insight_report.xlsx")
    wb.save(xlsx_path)
    print(f"Excel 报告已保存 -> {xlsx_path}")
    return xlsx_path


def main() -> None:
    ap = argparse.ArgumentParser(description="SoundInsight 一键洞察 Agent")
    ap.add_argument("--csv", required=True, help="待分析评论 CSV 路径")
    ap.add_argument("--format", choices=["md", "excel"], default="md",
                    help="报告格式：md 或 excel")
    ap.add_argument("--lang", choices=["zh", "en"], default="zh",
                    help="报告语言：zh 中文（默认）或 en 英文")
    args = ap.parse_args()
    report = analyze(args.csv, export=(args.format == "excel"),
                     lang=args.lang)
    print("\n" + report)


if __name__ == "__main__":
    main()
