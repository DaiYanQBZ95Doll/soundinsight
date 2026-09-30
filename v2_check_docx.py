# -*- coding: utf-8 -*-
"""抽查决赛主文档：表格填写是否完整、附录是否齐备。"""
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
DOC = "更新世界的锋芒_SoundInsight_决赛入围定稿作品.docx"
doc = Document(DOC)
print(f"段落 {len(doc.paragraphs)}｜表格 {len(doc.tables)}")
for ti, t in enumerate(doc.tables, 1):
    head = " / ".join(c.text.strip()[:14] for c in t.rows[0].cells)
    print(f"\n[表 {ti}] 表头：{head}｜行数 {len(t.rows)}")
    for r in t.rows[1:6]:
        cells = [c.text.strip().replace("\n", " ")[:34] for c in r.cells]
        if any(cells):
            print("    " + " | ".join(cells))
kind = {"附录 A": 0, "附录 B": 0, "附录 C": 0, "附录 D": 0, "附录 E": 0, "附录 F": 0}
for p in doc.paragraphs:
    for k in kind:
        if p.text.strip().startswith(k):
            kind[k] += 1
print("\n附录标题命中：", kind)
txt = "\n".join(p.text for p in doc.paragraphs)
for probe in ("更新世界的锋芒", "19195907942", "0.7220", "68.3%", "置信度档位", "已知表述勘误"):
    print(f"  {'OK  ' if probe in txt else 'MISS'} 含 {probe}")
