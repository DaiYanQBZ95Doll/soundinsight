# -*- coding: utf-8 -*-
# verify_docx.py —— 校验生成的 Word 主文档内容完整性
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

DOCX = sys.argv[1] if len(sys.argv) > 1 else \
    "更新世界的锋芒_SoundInsight_复赛作品.docx"

d = Document(DOCX)
print(f"段落总数: {len(d.paragraphs)} ｜ 表格: {len(d.tables)}")

print("--- 标题清单（前 14 个）:")
c = 0
for p in d.paragraphs:
    if p.style.name.startswith("Heading"):
        print(f"  {p.style.name} | {p.text[:60]}")
        c += 1
        if c >= 14:
            break

print("--- 第 1 个表格（团队信息）前两行:")
t = d.tables[0]
for r in t.rows[:2]:
    print("  " + " | ".join(x.text for x in r.cells))

body = "\n".join(p.text for p in d.paragraphs)
cells = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
txt = body + "\n" + cells
print("--- 关键词核对:")
for k in ["个人团队", "个人开发者", "token-plan.cn-beijing.maas.aliyuncs.com",
          "在线链接填写表", "modelscope.cn/studios", "0.6871",
          "未开展用户验证", "更新世界的锋芒"]:
    print(("  OK   " if k in txt else "  MISS ") + k)
